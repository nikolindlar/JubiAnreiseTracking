import io
import json
from datetime import date

import openpyxl
import pytest

from anreise import essensplan
from test_app import admin_login, app, client  # noqa: F401  (Fixtures)

WD = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]


def make_xlsx(year=2026, month=11, month_label="November", values=None, weekday_shift=0,
              lunch_label="Mittagesssen", n_days=None):
    """Nachbau der Vorlage „Dienstplan Hauswirtschaft“."""
    import calendar
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Dienstplan Hauswirtschaft"
    ws["G1"] = "Monat:"
    ws.merge_cells("G1:I1")
    ws["J1"] = month_label
    ws["W1"] = "Jahr:"
    ws.merge_cells("W1:X1")
    ws["Y1"] = year
    for label, row in (("Abreise", 6), ("Frühstück", 7), (lunch_label, 8), ("Abendessen", 9)):
        ws.cell(row, 1, label)
    n_days = n_days or calendar.monthrange(year, month)[1]
    for d in range(1, n_days + 1):
        ws.cell(3, d + 1, d)
        ws.cell(4, d + 1, WD[(date(year, month, d).weekday() + weekday_shift) % 7])
    for (row, d), v in (values or {}).items():
        ws.cell(row, d + 1, v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


NOV = {(6, 1): 28, (7, 1): 55, (9, 1): 57,
       (7, 2): 57, (9, 2): 30,
       (7, 6): 43, (8, 6): 13, (9, 6): 43,
       (8, 9): 95, (9, 9): 95}


def test_parse_template():
    p = essensplan.parse(make_xlsx(values=NOV))
    assert (p["year"], p["month"], len(p["days"])) == (2026, 11, 30)
    assert p["days"][0] == {"day": "2026-11-01", "breakfast": 55, "lunch": 0, "dinner": 57, "departures": 28}
    assert p["days"][5]["lunch"] == 13
    s = essensplan.summary(p["days"])
    # Lunchpakete = Frühstück − Mittagessen: 55 + 57 + (43 − 13); Mittag ohne Frühstück gibt keine
    assert s["packs"] == 55 + 57 + 30
    assert s["meals"] == s["breakfast"] + s["packs"] + s["lunch"] + s["dinner"]
    # Abendessen am 1. (57) = Frühstück am 2. (57) → keine Warnung; am 2. nicht
    assert any(w.startswith("2.:") for w in p["warnings"])
    assert not any(w.startswith("1.:") for w in p["warnings"])


def test_lunch_packs():
    assert essensplan.lunch_packs(43, 13) == 30
    assert essensplan.lunch_packs(0, 95) == 0


@pytest.mark.parametrize("kwargs, message", [
    ({"weekday_shift": 1}, "Wochentag passt nicht"),
    ({"month_label": "Novembr"}, "Monat „Novembr“ nicht erkannt"),
    ({"lunch_label": "Essen"}, "Mittagessen“ nicht gefunden"),
    ({"n_days": 29}, "fehlen die Tage 30"),
    ({"values": {(7, 3): "viele"}}, "Zelle D7"),
    ({"values": {(7, 3): -1}}, "keine ganze Zahl"),
])
def test_parse_errors(kwargs, message):
    with pytest.raises(essensplan.PlanError, match=message):
        essensplan.parse(make_xlsx(**kwargs))


def test_parse_not_excel():
    with pytest.raises(essensplan.PlanError, match="keine lesbare Excel"):
        essensplan.parse(io.BytesIO(b"kein excel"))


def upload(client, csrf, buf, name="Dienstplan_Hauswirtschaft_2611.xlsx"):
    return client.post("/verwaltung/essensplan", data={"csrf": csrf, "plan": (buf, name)},
                       content_type="multipart/form-data")


def confirm_from_preview(client, csrf, page):
    import html
    import re
    data = html.unescape(re.search(r'name="data" value="([^"]*)"', page).group(1))
    return client.post("/verwaltung/essensplan/uebernehmen",
                       data={"csrf": csrf, "data": data, "filename": "plan.xlsx"})


def test_upload_preview_confirm_and_dashboard(app, client):
    csrf = admin_login(client)
    page = upload(client, csrf, make_xlsx(values=NOV)).get_data(as_text=True)
    assert "November 2026 übernehmen" in page and "bereits importiert" not in page
    # Vorschau speichert noch nichts
    app.state["today"] = date(2026, 11, 6)
    before = client.get("/api/dashboard/verpflegung").get_json()
    assert before["plan_months"] == [] and before["today"]["estimated"] is True

    assert confirm_from_preview(client, csrf, page).status_code == 302
    d = client.get("/api/dashboard/verpflegung").get_json()
    assert d["plan_months"] == [11]
    assert d["today"] == {"breakfast": 43, "packs": 30, "lunch": 13, "dinner": 43, "estimated": False}
    # Jan–Okt pauschal (je Gast Frühstück, Lunchpaket, Abendessen), November nach Plan bis zum 6.
    per_day = 20000 / 334
    pauschal = per_day * 304
    nov = {"breakfast": 55 + 57 + 43, "packs": 55 + 57 + 30, "lunch": 13, "dinner": 57 + 30 + 43}
    assert d["meals_to_date"]["breakfast"] == pytest.approx(pauschal + nov["breakfast"])
    assert d["meals_to_date"]["lunch"] == 13
    equiv = pauschal * 1.0 + nov["breakfast"] * .25 + (nov["packs"] + nov["lunch"]) * .40 + nov["dinner"] * .35
    assert d["days_to_date"] == pytest.approx(equiv)
    assert d["veg_kg"] == pytest.approx(equiv * 3.81)
    assert d["estimated_to_date"] is True
    # Tage im Planmonat ohne Essen: keine Gäste (nicht pauschal)
    app.state["today"] = date(2026, 11, 3)
    assert client.get("/api/dashboard/verpflegung").get_json()["today"] is None

    admin = client.get("/verwaltung").get_data(as_text=True)
    assert "November 2026" in admin and "plan.xlsx" in admin

    # Erneuter Upload ersetzt den Monat vollständig
    page = upload(client, csrf, make_xlsx(values={(7, 6): 10})).get_data(as_text=True)
    assert "bereits importiert" in page
    confirm_from_preview(client, csrf, page)
    app.state["today"] = date(2026, 11, 6)
    d = client.get("/api/dashboard/verpflegung").get_json()
    assert d["today"] == {"breakfast": 10, "packs": 10, "lunch": 0, "dinner": 0, "estimated": False}
    assert d["meals_to_date"]["lunch"] == 0

    # Löschen → wieder pauschal
    assert client.post("/verwaltung/essensplan/2026-11/loeschen", data={"csrf": csrf}).status_code == 302
    d = client.get("/api/dashboard/verpflegung").get_json()
    assert d["plan_months"] == [] and d["today"]["estimated"] is True


def test_upload_error_shown(client):
    csrf = admin_login(client)
    page = upload(client, csrf, make_xlsx(weekday_shift=2)).get_data(as_text=True)
    assert "konnte nicht übernommen werden" in page and "übernehmen</button>" not in page
    page = upload(client, csrf, io.BytesIO(b"x"), name="plan.xls").get_data(as_text=True)
    assert "keine lesbare Excel" in page


def test_upload_requires_admin_and_valid_data(app, client):
    other = app.test_client()
    r = other.post("/verwaltung/essensplan", data={"plan": (make_xlsx(), "p.xlsx")},
                   content_type="multipart/form-data")
    assert r.status_code == 302 and "/verwaltung/login" in r.headers["Location"]
    csrf = admin_login(client)
    bad = {"year": 2026, "month": 11, "days": [{"day": "2026-11-01", "breakfast": -1}]}
    r = client.post("/verwaltung/essensplan/uebernehmen", data={"csrf": csrf, "data": json.dumps(bad)})
    assert r.status_code == 400
    r = client.post("/verwaltung/essensplan/uebernehmen", data={"csrf": csrf, "data": "{"})
    assert r.status_code == 400
    assert client.post("/verwaltung/essensplan/kaputt/loeschen", data={"csrf": csrf}).status_code == 404


def test_meal_weights_setting(app, client):
    csrf = admin_login(client)
    client.post("/verwaltung/einstellungen", data={
        "csrf": csrf, "meals_days_per_year": "20000", "meals_organic_share": "50",
        "meals_veg_kg_per_day": "3.81", "meals_mixed_kg_per_day": "5.63", "closed_12": "on",
        "screen_rotation_seconds": "30", "meals_weight_breakfast": "20",
        "meals_weight_lunch": "40", "meals_weight_dinner": "40"})
    d = client.get("/api/dashboard/verpflegung").get_json()
    assert d["weights"] == {"breakfast": 20, "lunch": 40, "dinner": 40, "packs": 40}
